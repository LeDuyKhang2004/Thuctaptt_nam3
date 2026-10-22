from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BUILD_DIR = Path(r"D:\Study\VKU\Nam2HK2\JavaFramework\Ontap\leave_management\.codex_tmp\word_report")
OUTPUT = Path(r"D:\Bao_cao_LeaveFlow_LeDuyKhang.docx")
LOGO = Path(r"D:\Study\VKU\Nam2HK2\JavaFramework\Ontap\leave_management\.codex_tmp\slide_english_report\template-inspect\assets\ppt\media\image2.png")

# Resolved design system: standard_business_brief.
# Named visual override: LeaveFlow teal is used only on the cover and database
# placeholders; body headings and tables retain the preset's blue/gray tokens.
FONT = "Calibri"
BLACK = "111111"
WHITE = "FFFFFF"
GRAY = "666666"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "D8DEE6"
HEADING_BLUE = "2E74B5"
HEADING_DARK_BLUE = "1F4D78"
INK_BLUE = "0B2545"
LEAVEFLOW_TEAL = "2FA7B1"
LEAVEFLOW_ORANGE = "F4A30A"
LIGHT_TEAL = "EAF7F8"

PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(
    run,
    *,
    name: str = FONT,
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color: str = "AAB4C2", size: int = 6) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa: Sequence[int], indent_dxa: int = TABLE_INDENT_DXA) -> None:
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must sum to {CONTENT_WIDTH_DXA}, got {sum(widths_dxa)}")

    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_styled_table(
    doc: Document,
    headers: Sequence[str],
    rows: Iterable[Sequence[str]],
    widths_dxa: Sequence[int],
    *,
    body_size: float = 9.2,
    center_columns: set[int] | None = None,
) -> object:
    rows = list(rows)
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    set_table_geometry(table, widths_dxa)
    center_columns = center_columns or set()

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        set_cell_shading(cell, LIGHT_GRAY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            set_run_font(run, size=9.3, color=INK_BLUE, bold=True)
    set_repeat_table_header(table.rows[0])

    for row_index, row_data in enumerate(rows, start=1):
        for col_index, value in enumerate(row_data):
            cell = table.rows[row_index].cells[col_index]
            cell.text = str(value)
            if row_index % 2 == 0:
                set_cell_shading(cell, "FAFBFC")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if col_index in center_columns else WD_ALIGN_PARAGRAPH.LEFT
            )
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            for run in paragraph.runs:
                set_run_font(run, size=body_size, color=BLACK)

    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.space_after = Pt(4)
    return table


def apply_left_border(paragraph, color: str = HEADING_BLUE, size: int = 18, space: int = 8) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), str(space))
    left.set(qn("w:color"), color)
    borders.append(left)


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_callout(doc: Document, label: str, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.left_indent = Inches(0.15)
    paragraph.paragraph_format.right_indent = Inches(0.15)
    paragraph.paragraph_format.line_spacing = 1.1
    apply_left_border(paragraph, color=HEADING_BLUE)
    shade_paragraph(paragraph, "F4F6F9")
    run = paragraph.add_run(f"{label}: ")
    set_run_font(run, size=10.5, color=HEADING_DARK_BLUE, bold=True)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5, color=BLACK)


def add_body(doc: Document, text: str, *, bold_lead: str | None = None) -> object:
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.widow_control = True
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_bullet(doc: Document, text: str, num_id: int, level: int = 0) -> object:
    paragraph = doc.add_paragraph(style="LF Bullet")
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_numbered(doc: Document, text: str, num_id: int, level: int = 0) -> object:
    paragraph = doc.add_paragraph(style="LF Number")
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), str(level))
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_numbering_definition(doc: Document, *, bullet: bool) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "hybridMultilevel")
    abstract.append(multi)

    for level in range(3):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), "•" if bullet else f"%{level + 1}.")
        lvl.append(lvl_text)
        suffix = OxmlElement("w:suff")
        suffix.set(qn("w:val"), "tab")
        lvl.append(suffix)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), str(720 + level * 360))
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(720 + level * 360))
        ind.set(qn("w:hanging"), "360")
        p_pr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "160")
        spacing.set(qn("w:line"), "280")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.append(spacing)
        lvl.append(p_pr)
        abstract.append(lvl)

    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_page_number(paragraph) -> None:
    paragraph.add_run("Trang ")
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, separate, text, end])


def add_toc_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "Mục lục được cập nhật tự động khi mở trong Microsoft Word."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, display, end])


def set_update_fields_on_open(doc: Document) -> None:
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.keep_with_next = False
    run = paragraph.add_run(text)
    set_run_font(run, size=9.5, color=GRAY, italic=True)


def make_placeholder(path: Path, title: str, subtitle: str) -> None:
    width, height = 1500, 650
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    regular_path = r"C:\Windows\Fonts\arial.ttf"
    bold_path = r"C:\Windows\Fonts\arialbd.ttf"
    font_title = ImageFont.truetype(bold_path, 48)
    font_subtitle = ImageFont.truetype(regular_path, 28)
    font_hint = ImageFont.truetype(regular_path, 23)

    # Border.
    dash, gap, margin = 34, 22, 34
    for x in range(margin, width - margin, dash + gap):
        draw.line((x, margin, min(x + dash, width - margin), margin), fill="#2FA7B1", width=5)
        draw.line((x, height - margin, min(x + dash, width - margin), height - margin), fill="#2FA7B1", width=5)
    for y in range(margin, height - margin, dash + gap):
        draw.line((margin, y, margin, min(y + dash, height - margin)), fill="#2FA7B1", width=5)
        draw.line((width - margin, y, width - margin, min(y + dash, height - margin)), fill="#2FA7B1", width=5)

    # Simple database cylinder icon.
    cx, top = width // 2, 82
    icon_w, icon_h = 220, 145
    left, right = cx - icon_w // 2, cx + icon_w // 2
    draw.ellipse((left, top, right, top + 55), fill="#EAF7F8", outline="#2FA7B1", width=7)
    draw.rectangle((left, top + 27, right, top + icon_h - 27), fill="#EAF7F8", outline="#2FA7B1", width=7)
    draw.ellipse((left, top + icon_h - 55, right, top + icon_h), fill="#EAF7F8", outline="#2FA7B1", width=7)
    draw.arc((left, top + 55, right, top + 110), start=0, end=180, fill="#2FA7B1", width=5)

    title_box = draw.textbbox((0, 0), title, font=font_title)
    draw.text(((width - (title_box[2] - title_box[0])) / 2, 285), title, font=font_title, fill="#1F4D78")
    subtitle_box = draw.textbbox((0, 0), subtitle, font=font_subtitle)
    draw.text(((width - (subtitle_box[2] - subtitle_box[0])) / 2, 365), subtitle, font=font_subtitle, fill="#333333")
    hint = "Nhấp chuột phải vào ảnh → Change Picture → From a File"
    hint_box = draw.textbbox((0, 0), hint, font=font_hint)
    draw.text(((width - (hint_box[2] - hint_box[0])) / 2, 530), hint, font=font_hint, fill="#666666")
    image.save(path)


def add_placeholder_figure(doc: Document, image_path: Path, alt_text: str, caption: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    inline_shape = run.add_picture(str(image_path), width=Inches(6.2))
    inline_shape._inline.docPr.set("descr", alt_text)
    inline_shape._inline.docPr.set("title", alt_text)
    add_caption(doc, caption)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(BLACK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    normal.paragraph_format.widow_control = True

    heading_tokens = {
        "Heading 1": (16, HEADING_BLUE, 16, 8, True),
        "Heading 2": (13, HEADING_BLUE, 12, 6, False),
        "Heading 3": (12, HEADING_DARK_BLUE, 8, 4, False),
    }
    for name, (size, color, before, after, page_break) in heading_tokens.items():
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.page_break_before = page_break

    if "Front Matter Heading" not in styles:
        style = styles.add_style("Front Matter Heading", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["Front Matter Heading"]
    style.font.name = FONT
    style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(18)
    style.font.bold = True
    style.font.color.rgb = rgb(INK_BLUE)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style.paragraph_format.space_before = Pt(8)
    style.paragraph_format.space_after = Pt(18)
    style.paragraph_format.keep_with_next = True

    if "Front Matter Subheading" not in styles:
        style = styles.add_style("Front Matter Subheading", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["Front Matter Subheading"]
    style.font.name = FONT
    style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(13)
    style.font.bold = True
    style.font.color.rgb = rgb(HEADING_BLUE)
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.keep_together = True

    for name in ("LF Bullet", "LF Number"):
        if name not in styles:
            style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(11)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167
        style.paragraph_format.widow_control = True

    caption = styles["Caption"]
    caption.font.name = FONT
    caption._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    caption.font.size = Pt(9.5)
    caption.font.italic = True
    caption.font.color.rgb = rgb(GRAY)


def configure_sections(doc: Document) -> None:
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
        section.different_first_page_header_footer = True

        header = section.header
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
        run = p.add_run("BÁO CÁO THỰC TẬP THỰC TẾ")
        set_run_font(run, size=8.5, color=GRAY, bold=True)
        run = p.add_run("\tLEAVEFLOW")
        set_run_font(run, size=8.5, color=GRAY, bold=True)

        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        add_page_number(p)
        for run in p.runs:
            set_run_font(run, size=8.5, color=GRAY)


def front_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Front Matter Heading")
    p.add_run(text)


def front_subheading(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Front Matter Subheading")
    p.add_run(text)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)


def add_manual_link_line(doc: Document, label: str, value: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    label_run = p.add_run(label)
    set_run_font(label_run, size=10.5, color=HEADING_DARK_BLUE, bold=True)
    value_run = p.add_run(value)
    set_run_font(value_run, size=10.5, color=BLACK)


def build_document() -> None:
    BUILD_DIR.mkdir(parents=True, exist_ok=True)
    placeholder_erd = BUILD_DIR / "db_placeholder_erd.png"
    placeholder_users = BUILD_DIR / "db_placeholder_users.png"
    placeholder_leaves = BUILD_DIR / "db_placeholder_leave_requests.png"
    make_placeholder(placeholder_erd, "CHÈN ẢNH DATABASE TẠI ĐÂY", "Sơ đồ ERD của hệ thống LeaveFlow")
    make_placeholder(placeholder_users, "CHÈN ẢNH DATABASE TẠI ĐÂY", "Cấu trúc hoặc dữ liệu mẫu của bảng users")
    make_placeholder(placeholder_leaves, "CHÈN ẢNH DATABASE TẠI ĐÂY", "Cấu trúc hoặc dữ liệu mẫu của bảng leave_requests")

    doc = Document()
    configure_styles(doc)
    configure_sections(doc)
    bullet_num_id = add_numbering_definition(doc, bullet=True)
    set_update_fields_on_open(doc)

    doc.core_properties.title = "Báo cáo thực tập thực tế - LeaveFlow"
    doc.core_properties.subject = "Hệ thống quản lý nghỉ phép và nhân sự tích hợp chatbot"
    doc.core_properties.author = "Lê Duy Khang"
    doc.core_properties.keywords = "LeaveFlow, Spring MVC, Java, MySQL, Chatbot, Leave Management"
    doc.core_properties.comments = "Báo cáo được xây dựng từ đề cương và mã nguồn LeaveFlow."

    # Cover: editorial_cover pattern adapted for an academic report.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    if LOGO.exists():
        shape = p.add_run().add_picture(str(LOGO), width=Inches(1.55))
        shape._inline.docPr.set("descr", "Logo Trường Đại học Công nghệ Thông tin và Truyền thông Việt - Hàn")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN\nVÀ TRUYỀN THÔNG VIỆT - HÀN")
    set_run_font(run, size=13, color=INK_BLUE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(26)
    run = p.add_run("KHOA KHOA HỌC MÁY TÍNH")
    set_run_font(run, size=11.5, color=HEADING_BLUE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("BÁO CÁO THỰC TẬP THỰC TẾ")
    set_run_font(run, size=24, color=INK_BLUE, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.right_indent = Inches(0.35)
    p.paragraph_format.space_after = Pt(28)
    run = p.add_run("GIẢI PHÁP TÍCH HỢP CHATBOT HỖ TRỢ TRA CỨU THÔNG TIN VÀO HỆ THỐNG QUẢN LÝ NGHỈ PHÉP VÀ NHÂN SỰ DOANH NGHIỆP")
    set_run_font(run, size=17, color=LEAVEFLOW_TEAL, bold=True)

    for label, value in (
        ("Sinh viên thực hiện", "LÊ DUY KHANG"),
        ("Mã số sinh viên", "24IT105"),
        ("Lớp", "24SE1"),
        ("Giảng viên hướng dẫn", "TS. Nguyễn Văn Bình"),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"{label}: ")
        set_run_font(r, size=11, color=GRAY, bold=True)
        r = p.add_run(value)
        set_run_font(r, size=11, color=BLACK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(34)
    run = p.add_run("Đà Nẵng, tháng 8 năm 2026")
    set_run_font(run, size=11, color=GRAY, italic=True)
    doc.add_page_break()

    # Report information.
    front_heading(doc, "THÔNG TIN BÁO CÁO")
    add_styled_table(
        doc,
        ("Nội dung", "Thông tin"),
        (
            ("Họ và tên sinh viên", "Lê Duy Khang"),
            ("MSSV / Lớp", "24IT105 / 24SE1"),
            ("Ngành", "Công nghệ Thông tin - Kỹ sư Công nghệ phần mềm"),
            ("Điện thoại", "0779598550"),
            ("Email", "khangld.24it@vku.udn.vn"),
            ("Giảng viên hướng dẫn", "TS. Nguyễn Văn Bình"),
            ("Tên hệ thống", "LeaveFlow - Hệ thống quản lý nghỉ phép và nhân sự tích hợp chatbot"),
            ("Thời gian thực hiện", "29/06/2026 - 17/08/2026"),
        ),
        (2700, 6660),
        body_size=10,
    )
    add_callout(
        doc,
        "Phạm vi báo cáo",
        "Tài liệu trình bày quá trình phân tích, thiết kế, cài đặt và kiểm thử hệ thống LeaveFlow dựa trên mã nguồn hiện tại. Các ảnh database trong Chương 4 được để sẵn dưới dạng ảnh placeholder để người thực hiện thay bằng ảnh chụp thực tế.",
    )
    front_subheading(doc, "Kế hoạch thực hiện")
    add_styled_table(
        doc,
        ("Thời gian", "Nội dung chính", "Kết quả"),
        (
            ("29/06 - 05/07", "Khảo sát quy trình nghỉ phép và nhu cầu số hóa", "Xác định bài toán và phạm vi"),
            ("06/07 - 10/07", "Nghiên cứu công nghệ, kiến trúc phân lớp", "Lựa chọn Spring MVC, JPA, MySQL"),
            ("11/07 - 18/07", "Phân tích yêu cầu và thiết kế hệ thống", "Xác định vai trò, dữ liệu và luồng xử lý"),
            ("19/07 - 26/07", "Xây dựng chức năng nhân viên", "Đăng ký, đăng nhập, tạo và theo dõi đơn"),
            ("27/07 - 02/08", "Xây dựng chức năng quản trị", "Duyệt đơn, lọc trạng thái, quản lý nhân viên"),
            ("03/08 - 09/08", "Tích hợp thống kê và chatbot", "Tra cứu theo vai trò và gợi ý thao tác"),
            ("10/08 - 14/08", "Kiểm thử và hoàn thiện", "12 kiểm thử tự động đạt yêu cầu"),
            ("15/08 - 17/08", "Hoàn thiện báo cáo và slide", "Tổng hợp sản phẩm bàn giao"),
        ),
        (1800, 3700, 3860),
        body_size=8.7,
        center_columns={0},
    )
    doc.add_page_break()

    # Acknowledgements.
    front_heading(doc, "LỜI CẢM ƠN")
    add_body(doc, "Trong quá trình thực hiện đề tài, em đã nhận được sự hướng dẫn, hỗ trợ và góp ý từ giảng viên, bạn bè và các thầy cô trong Khoa Khoa học Máy tính. Những góp ý này giúp em hoàn thiện tư duy phân tích hệ thống, củng cố kiến thức Java Web và rèn luyện kỹ năng tổ chức một dự án phần mềm theo mô hình phân lớp.")
    add_body(doc, "Em xin gửi lời cảm ơn chân thành đến TS. Nguyễn Văn Bình đã định hướng nội dung, theo dõi tiến độ và góp ý cho quá trình xây dựng hệ thống LeaveFlow. Sự hướng dẫn của thầy giúp em nhận ra các vấn đề thực tế trong quản lý nghỉ phép, đồng thời biết cách chuyển hóa yêu cầu nghiệp vụ thành các chức năng có thể triển khai và kiểm thử.")
    add_body(doc, "Em cũng xin cảm ơn Trường Đại học Công nghệ Thông tin và Truyền thông Việt - Hàn cùng Khoa Khoa học Máy tính đã tạo môi trường học tập và thực hành thuận lợi. Kiến thức về lập trình Java, cơ sở dữ liệu, mô hình MVC và kiểm thử phần mềm là nền tảng trực tiếp để em hoàn thành đề tài này.")
    add_body(doc, "Do thời gian và kinh nghiệm thực tế còn hạn chế, báo cáo khó tránh khỏi thiếu sót. Em mong tiếp tục nhận được các ý kiến đóng góp để có thể cải thiện hệ thống về tính bảo mật, khả năng mở rộng và chất lượng trải nghiệm người dùng.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(28)
    r = p.add_run("Sinh viên thực hiện\nLê Duy Khang")
    set_run_font(r, size=11, color=BLACK, bold=True)
    doc.add_page_break()

    # Abstract.
    front_heading(doc, "TÓM TẮT")
    add_body(doc, "LeaveFlow là ứng dụng web quản lý nghỉ phép và nhân sự cơ bản dành cho hai nhóm người dùng: nhân viên và quản trị viên. Hệ thống số hóa quy trình từ lúc nhân viên tạo đơn, theo dõi trạng thái đến khi quản trị viên duyệt hoặc từ chối. Dữ liệu được lưu trong MySQL và truy cập thông qua lớp repository của Spring Data JPA.")
    add_body(doc, "Điểm mở rộng của đề tài là chatbot hỗ trợ tra cứu theo ngữ cảnh đăng nhập. Trợ lý có thể giải thích trạng thái PENDING, APPROVED và REJECTED; tra cứu đơn theo mã; hiển thị đơn gần nhất; thống kê số lượng đơn; tính tổng ngày nghỉ đã được duyệt; và gợi ý thao tác tiếp theo. Phạm vi dữ liệu được giới hạn theo vai trò, do đó nhân viên chỉ nhận thông tin thuộc tài khoản của mình, còn quản trị viên có thể xem số liệu toàn hệ thống.")
    add_body(doc, "Sản phẩm được xây dựng với Java 17, Spring MVC 6.1, JSP/JSTL, Spring Data JPA, Hibernate 6.5 và MySQL 8.4, đóng gói dưới dạng WAR để triển khai trên Apache Tomcat. Mật khẩu được bảo vệ bằng PBKDF2WithHmacSHA256 với salt ngẫu nhiên và 210.000 vòng lặp. Bộ kiểm thử hiện có gồm 12 trường hợp cho dịch vụ mật khẩu và chatbot; tất cả đều đạt.")
    add_callout(doc, "Từ khóa", "LeaveFlow; quản lý nghỉ phép; quản lý nhân sự; chatbot; Spring MVC; JPA; Hibernate; MySQL; PBKDF2.")
    doc.add_page_break()

    # TOC.
    front_heading(doc, "MỤC LỤC")
    toc = doc.add_paragraph()
    toc.paragraph_format.space_after = Pt(6)
    add_toc_field(toc)
    doc.add_page_break()

    # Lists.
    front_heading(doc, "DANH MỤC HÌNH, BẢNG VÀ TỪ VIẾT TẮT")
    front_subheading(doc, "Danh mục hình")
    for text in (
        "Hình 4.1. Sơ đồ quan hệ cơ sở dữ liệu của hệ thống LeaveFlow.",
        "Hình 4.2. Cấu trúc hoặc dữ liệu mẫu của bảng users.",
        "Hình 4.3. Cấu trúc hoặc dữ liệu mẫu của bảng leave_requests.",
    ):
        add_bullet(doc, text, bullet_num_id)
    front_subheading(doc, "Danh mục bảng")
    for text in (
        "Bảng 1.1. Mục tiêu cụ thể của đề tài.",
        "Bảng 2.1. Yêu cầu chức năng theo vai trò.",
        "Bảng 3.1. Công nghệ sử dụng trong hệ thống.",
        "Bảng 4.1. Cấu trúc bảng users.",
        "Bảng 4.2. Cấu trúc bảng leave_requests.",
        "Bảng 6.1. Các trường hợp kiểm thử tiêu biểu.",
    ):
        add_bullet(doc, text, bullet_num_id)
    front_subheading(doc, "Từ viết tắt")
    add_styled_table(
        doc,
        ("Từ viết tắt", "Diễn giải"),
        (
            ("MVC", "Model - View - Controller"),
            ("JPA", "Jakarta Persistence API"),
            ("ORM", "Object - Relational Mapping"),
            ("JSP", "Jakarta Server Pages"),
            ("JSTL", "Jakarta Standard Tag Library"),
            ("PBKDF2", "Password-Based Key Derivation Function 2"),
            ("ERD", "Entity Relationship Diagram"),
            ("WAR", "Web Application Archive"),
        ),
        (2100, 7260),
        body_size=10,
        center_columns={0},
    )

    # Chapter 1.
    add_heading(doc, "CHƯƠNG 1. TỔNG QUAN ĐỀ TÀI", 1)
    add_heading(doc, "1.1. Bối cảnh và lý do chọn đề tài", 2)
    add_body(doc, "Trong nhiều doanh nghiệp nhỏ và nhóm dự án, quy trình xin nghỉ phép vẫn được thực hiện qua tin nhắn, giấy tờ hoặc bảng tính rời rạc. Cách làm này khiến thông tin dễ bị thất lạc, khó xác định trách nhiệm phê duyệt, thiếu lịch sử xử lý và không thuận tiện khi cần tổng hợp số liệu.")
    add_body(doc, "Việc xây dựng một hệ thống web tập trung giúp nhân viên chủ động tạo và theo dõi đơn, đồng thời cung cấp cho quản trị viên một không gian thống nhất để xem, lọc và xử lý yêu cầu. Chatbot được tích hợp nhằm rút ngắn thời gian tìm kiếm thông tin và hướng dẫn người dùng ngay tại màn hình đang thao tác.")
    add_callout(doc, "Bài toán trung tâm", "Xây dựng một quy trình nghỉ phép rõ ràng, có trạng thái, có kiểm soát quyền truy cập và có trợ lý hội thoại hỗ trợ tra cứu nhanh.")
    add_heading(doc, "1.2. Mục tiêu", 2)
    add_styled_table(
        doc,
        ("Nhóm mục tiêu", "Nội dung"),
        (
            ("Nghiệp vụ", "Số hóa quá trình tạo, theo dõi, duyệt, từ chối và hủy đơn nghỉ phép."),
            ("Quản trị", "Cung cấp thống kê trạng thái, tìm kiếm nhân viên và cập nhật vai trò."),
            ("Trải nghiệm", "Thiết kế giao diện dễ sử dụng cho hai nhóm vai trò."),
            ("Chatbot", "Hỗ trợ tra cứu đơn, thống kê, giải thích trạng thái và gợi ý bước tiếp theo."),
            ("Kỹ thuật", "Áp dụng Java Web, Spring MVC, JPA/Hibernate, MySQL và kiểm thử tự động."),
        ),
        (2100, 7260),
        body_size=9.6,
    )
    add_caption(doc, "Bảng 1.1. Mục tiêu cụ thể của đề tài.")
    add_heading(doc, "1.3. Phạm vi thực hiện", 2)
    for item in (
        "Quản lý tài khoản với hai vai trò EMPLOYEE và ADMIN.",
        "Quản lý vòng đời đơn nghỉ phép với ba trạng thái PENDING, APPROVED và REJECTED.",
        "Tra cứu và thống kê dữ liệu ở mức cơ bản, chưa triển khai hạn mức phép năm hoặc chính sách nội bộ phức tạp.",
        "Chatbot xử lý theo luật và dữ liệu hệ thống hiện có; phiên bản hiện tại chưa gọi mô hình ngôn ngữ lớn bên ngoài.",
        "Triển khai thử nghiệm cục bộ với MySQL và Apache Tomcat.",
    ):
        add_bullet(doc, item, bullet_num_id)
    doc.add_page_break()
    add_heading(doc, "1.4. Phương pháp thực hiện", 2)
    chapter1_num_id = add_numbering_definition(doc, bullet=False)
    for item in (
        "Khảo sát quy trình nghiệp vụ và xác định vấn đề cần giải quyết.",
        "Phân tích vai trò, chức năng, dữ liệu và các quy tắc truy cập.",
        "Thiết kế kiến trúc phân lớp và mô hình cơ sở dữ liệu.",
        "Cài đặt theo từng nhóm chức năng, sau đó tích hợp chatbot.",
        "Kiểm thử dịch vụ quan trọng và rà soát toàn bộ luồng người dùng.",
    ):
        add_numbered(doc, item, chapter1_num_id)
    add_heading(doc, "1.5. Sản phẩm bàn giao", 2)
    for item in (
        "Mã nguồn ứng dụng được tổ chức theo dự án Maven và đóng gói WAR.",
        "Mô hình dữ liệu MySQL gồm hai thực thể chính, được ánh xạ bằng JPA/Hibernate.",
        "Bộ kiểm thử tự động gồm 12 trường hợp cho PasswordService và ChatbotService.",
        "Báo cáo Word và slide thuyết trình tổng hợp kết quả thực hiện đề tài.",
    ):
        add_bullet(doc, item, bullet_num_id)
    add_heading(doc, "1.6. Cấu trúc báo cáo", 2)
    add_body(doc, "Báo cáo gồm bảy chương. Chương 1 giới thiệu đề tài; Chương 2 phân tích yêu cầu; Chương 3 trình bày kiến trúc và công nghệ; Chương 4 mô tả cơ sở dữ liệu; Chương 5 trình bày quá trình cài đặt chức năng; Chương 6 tổng hợp kiểm thử và đánh giá; Chương 7 đề xuất hướng phát triển. Phần cuối gồm kết luận, tài liệu tham khảo và hướng dẫn thay ảnh database.")

    # Chapter 2.
    add_heading(doc, "CHƯƠNG 2. PHÂN TÍCH VÀ YÊU CẦU HỆ THỐNG", 1)
    add_heading(doc, "2.1. Tác nhân và quyền hạn", 2)
    add_body(doc, "LeaveFlow phân chia người dùng thành hai vai trò chính. Nhân viên thao tác trên dữ liệu cá nhân, trong khi quản trị viên xử lý dữ liệu toàn hệ thống. Việc phân quyền được kiểm tra tại controller và tiếp tục được áp dụng trong các nghiệp vụ tra cứu của chatbot.")
    add_styled_table(
        doc,
        ("Vai trò", "Chức năng chính", "Phạm vi dữ liệu"),
        (
            ("EMPLOYEE", "Đăng ký, đăng nhập, tạo đơn, xem danh sách và chi tiết, hủy đơn PENDING, dùng chatbot", "Chỉ dữ liệu của tài khoản hiện tại"),
            ("ADMIN", "Xem dashboard, lọc đơn, duyệt/từ chối, tìm kiếm nhân viên, cập nhật vai trò, dùng chatbot", "Toàn bộ dữ liệu hệ thống"),
        ),
        (1500, 4300, 3560),
        body_size=9.3,
    )
    add_caption(doc, "Bảng 2.1. Yêu cầu chức năng theo vai trò.")
    add_heading(doc, "2.2. Yêu cầu chức năng", 2)
    add_heading(doc, "2.2.1. Nhóm chức năng xác thực", 3)
    for item in (
        "Đăng ký tài khoản mới với họ tên, username, email và mật khẩu tối thiểu tám ký tự.",
        "Kiểm tra username trùng lặp trước khi lưu.",
        "Đăng nhập và điều hướng theo vai trò; đăng xuất bằng cách hủy session.",
        "Tự động gán vai trò EMPLOYEE cho tài khoản mới.",
    ):
        add_bullet(doc, item, bullet_num_id)
    add_heading(doc, "2.2.2. Nhóm chức năng đơn nghỉ phép", 3)
    for item in (
        "Nhân viên tạo đơn với lý do, ngày bắt đầu và ngày kết thúc.",
        "Hệ thống kiểm tra lý do tối thiểu năm ký tự và ngày kết thúc không trước ngày bắt đầu.",
        "Đơn mới được gán trạng thái PENDING và thời điểm tạo tự động.",
        "Nhân viên chỉ được xem hoặc hủy đơn của mình; chỉ đơn PENDING mới được hủy.",
        "Quản trị viên duyệt hoặc từ chối đơn và lưu ghi chú xử lý.",
    ):
        add_bullet(doc, item, bullet_num_id)
    add_heading(doc, "2.2.3. Nhóm chức năng quản trị nhân sự", 3)
    for item in (
        "Tìm kiếm nhân viên theo họ tên, username hoặc email, không phân biệt hoa thường.",
        "Cập nhật vai trò giữa EMPLOYEE và ADMIN; không cho phép quản trị viên tự hạ quyền của chính mình.",
        "Thống kê tổng số đơn và số đơn theo từng trạng thái trên dashboard.",
    ):
        add_bullet(doc, item, bullet_num_id)
    add_heading(doc, "2.3. Yêu cầu phi chức năng", 2)
    for item in (
        "Bảo mật: mật khẩu không lưu ở dạng rõ; dữ liệu chatbot được giới hạn theo vai trò.",
        "Dễ sử dụng: các trang có cấu trúc thống nhất, trạng thái hiển thị rõ và chatbot xuất hiện xuyên suốt phiên làm việc.",
        "Dễ bảo trì: mã nguồn được tách thành controller, service, repository, entity và view.",
        "Khả năng triển khai: dự án đóng gói WAR, có thể chạy trên Tomcat và kết nối MySQL bằng cấu hình thuộc tính.",
        "Khả năng kiểm thử: nghiệp vụ mật khẩu và chatbot được tách thành service để kiểm thử độc lập.",
    ):
        add_bullet(doc, item, bullet_num_id)
    add_heading(doc, "2.4. Luồng xử lý chính", 2)
    chapter2_num_id = add_numbering_definition(doc, bullet=False)
    for item in (
        "Nhân viên đăng nhập và nhập thông tin đơn nghỉ phép.",
        "Controller kiểm tra dữ liệu; service gắn người tạo, trạng thái PENDING và thời điểm tạo.",
        "Quản trị viên mở danh sách chờ duyệt, xem lý do và khoảng thời gian nghỉ.",
        "Quản trị viên chọn APPROVED hoặc REJECTED và có thể ghi chú.",
        "Nhân viên theo dõi kết quả trong danh sách cá nhân hoặc hỏi chatbot.",
    ):
        add_numbered(doc, item, chapter2_num_id)

    # Chapter 3.
    add_heading(doc, "CHƯƠNG 3. KIẾN TRÚC VÀ CÔNG NGHỆ", 1)
    add_heading(doc, "3.1. Kiến trúc tổng thể", 2)
    add_body(doc, "Hệ thống sử dụng kiến trúc MVC kết hợp phân lớp. DispatcherServlet tiếp nhận request và chuyển đến controller phù hợp. Controller thực hiện kiểm tra session và dữ liệu đầu vào, sau đó gọi service để xử lý nghiệp vụ. Repository cung cấp các thao tác truy vấn JPA, entity ánh xạ với bảng MySQL, còn JSP/JSTL chịu trách nhiệm hiển thị giao diện.")
    add_callout(doc, "Luồng kỹ thuật", "Browser → DispatcherServlet → Controller → Service → Repository/JPA → MySQL; kết quả được đưa trở lại JSP để hiển thị.")
    add_heading(doc, "3.2. Công nghệ sử dụng", 2)
    add_styled_table(
        doc,
        ("Lớp", "Công nghệ", "Vai trò"),
        (
            ("Giao diện", "JSP, JSTL, HTML5, CSS3, JavaScript", "Hiển thị trang, biểu mẫu, trạng thái và chatbot widget"),
            ("Ứng dụng", "Java 17, Spring MVC 6.1.14", "Điều hướng request, validation và điều phối nghiệp vụ"),
            ("Nghiệp vụ", "Service components", "Quản lý người dùng, đơn nghỉ phép, mật khẩu và chatbot"),
            ("Dữ liệu", "Spring Data JPA 3.3.5, Hibernate 6.5.3", "Ánh xạ entity và tạo repository"),
            ("Cơ sở dữ liệu", "MySQL 8.4", "Lưu tài khoản và đơn nghỉ phép"),
            ("Triển khai", "Maven WAR, Apache Tomcat", "Đóng gói và chạy ứng dụng web"),
            ("Kiểm thử", "JUnit Jupiter 5.11.3", "Kiểm thử dịch vụ mật khẩu và chatbot"),
        ),
        (1500, 3100, 4760),
        body_size=9.1,
    )
    add_caption(doc, "Bảng 3.1. Công nghệ sử dụng trong hệ thống.")
    add_heading(doc, "3.3. Tổ chức mã nguồn", 2)
    for item in (
        "com.javaweb.controller: nhận request, kiểm tra quyền và lựa chọn view.",
        "com.javaweb.service: thực hiện các quy tắc nghiệp vụ và bảo vệ phạm vi dữ liệu.",
        "com.javaweb.repository: kế thừa JpaRepository và định nghĩa truy vấn theo quy ước tên hàm.",
        "com.javaweb.entity: ánh xạ bảng users và leave_requests.",
        "WEB-INF/views: chứa JSP cho trang xác thực, nhân viên, quản trị viên và component dùng chung.",
        "assets: chứa CSS và JavaScript của giao diện, bao gồm widget chatbot.",
    ):
        add_bullet(doc, item, bullet_num_id)
    add_heading(doc, "3.4. Bảo mật và kiểm soát truy cập", 2)
    add_body(doc, "Mật khẩu được băm bằng PBKDF2WithHmacSHA256, sử dụng salt ngẫu nhiên 16 byte, khóa 256 bit và 210.000 vòng lặp. Chuỗi lưu trữ gồm tiền tố, số vòng lặp, salt và giá trị dẫn xuất được mã hóa Base64. Khi đăng nhập, phép so sánh dùng MessageDigest.isEqual để giảm rủi ro phân tích thời gian.")
    add_body(doc, "Ứng dụng hiện sử dụng HttpSession để lưu currentUser và kiểm tra vai trò tại controller. Chatbot tiếp tục áp dụng phạm vi dữ liệu: nhân viên chỉ được tra cứu đơn của mình, còn quản trị viên được phép xem dữ liệu toàn hệ thống.")
    add_callout(doc, "Lưu ý triển khai", "Cấu hình hibernate.hbm2ddl.auto=update, tài khoản MySQL root không mật khẩu và cơ chế kiểm tra quyền thủ công chỉ phù hợp môi trường học tập. Khi triển khai thật cần dùng biến môi trường, connection pool, HTTPS, Spring Security và CSRF protection.")
    add_heading(doc, "3.5. Cấu hình hệ thống", 2)
    for item in (
        "Property placeholder đọc thông tin kết nối JDBC và các thuộc tính Hibernate từ tệp cấu hình.",
        "Component scan và MVC annotation kích hoạt controller, service và xử lý tài nguyên tĩnh.",
        "EntityManagerFactory, JPA repository và transaction manager đảm nhiệm truy cập dữ liệu có giao dịch.",
        "View resolver ánh xạ tên view đến các tệp JSP nằm trong WEB-INF/views.",
    ):
        add_bullet(doc, item, bullet_num_id)
    add_heading(doc, "3.6. Đóng gói và triển khai", 2)
    chapter3_num_id = add_numbering_definition(doc, bullet=False)
    for item in (
        "Biên dịch và kiểm thử dự án bằng Maven với môi trường Java 17.",
        "Đóng gói ứng dụng thành tệp leave_management.war.",
        "Triển khai WAR trên Apache Tomcat và kết nối đến database MySQL.",
        "Kiểm tra lại đăng nhập, luồng đơn nghỉ phép, phân quyền và chatbot sau khi triển khai.",
    ):
        add_numbered(doc, item, chapter3_num_id)

    # Chapter 4 - Database placeholders.
    add_heading(doc, "CHƯƠNG 4. THIẾT KẾ CƠ SỞ DỮ LIỆU", 1)
    add_heading(doc, "4.1. Mô hình dữ liệu", 2)
    add_body(doc, "Cơ sở dữ liệu của LeaveFlow tập trung vào hai thực thể chính: người dùng và đơn nghỉ phép. Mỗi người dùng có thể tạo nhiều đơn; mỗi đơn chỉ thuộc về một người dùng. Khóa ngoại user_id trong leave_requests đảm bảo liên kết giữa hai bảng.")
    add_callout(doc, "Cách thay ảnh", "Nhấp chuột phải vào ảnh placeholder, chọn Change Picture → From a File, sau đó chọn ảnh database của bạn. Khung ảnh đã được đặt đúng kích thước nên bố cục và chú thích sẽ được giữ nguyên.")
    add_placeholder_figure(doc, placeholder_erd, "Vị trí chèn ảnh sơ đồ ERD của LeaveFlow", "Hình 4.1. Sơ đồ quan hệ cơ sở dữ liệu của hệ thống LeaveFlow.")

    add_heading(doc, "4.2. Bảng users", 2)
    add_body(doc, "Bảng users lưu thông tin tài khoản và vai trò. Trường username là duy nhất; password lưu chuỗi PBKDF2; role nhận giá trị EMPLOYEE hoặc ADMIN. Quan hệ một - nhiều với leave_requests được khai báo bằng @OneToMany trong UserEntity.")
    add_styled_table(
        doc,
        ("Trường", "Kiểu dữ liệu", "Ràng buộc", "Ý nghĩa"),
        (
            ("id", "BIGINT", "PK, AUTO_INCREMENT", "Mã người dùng"),
            ("username", "VARCHAR", "UNIQUE, NOT NULL", "Tên đăng nhập"),
            ("password", "VARCHAR(255)", "NOT NULL", "Mật khẩu đã băm"),
            ("email", "VARCHAR", "Có thể rỗng", "Email liên hệ"),
            ("full_name", "VARCHAR", "Có thể rỗng", "Họ và tên"),
            ("role", "VARCHAR", "EMPLOYEE / ADMIN", "Vai trò truy cập"),
        ),
        (1400, 1800, 2600, 3560),
        body_size=9.0,
        center_columns={0, 1},
    )
    add_caption(doc, "Bảng 4.1. Cấu trúc bảng users.")
    add_placeholder_figure(doc, placeholder_users, "Vị trí chèn ảnh cấu trúc hoặc dữ liệu bảng users", "Hình 4.2. Cấu trúc hoặc dữ liệu mẫu của bảng users.")

    add_heading(doc, "4.3. Bảng leave_requests", 2)
    add_body(doc, "Bảng leave_requests lưu nội dung đơn, khoảng thời gian nghỉ, trạng thái, ghi chú quản trị viên và thời điểm tạo. Trường reason và admin_note dùng kiểu TEXT để hỗ trợ nội dung dài; start_date và end_date dùng LocalDate trong Java; created_at dùng LocalDateTime.")
    add_styled_table(
        doc,
        ("Trường", "Kiểu dữ liệu", "Ràng buộc", "Ý nghĩa"),
        (
            ("id", "BIGINT", "PK, AUTO_INCREMENT", "Mã đơn nghỉ phép"),
            ("reason", "TEXT", "Tối thiểu 5 ký tự ở tầng ứng dụng", "Lý do nghỉ"),
            ("start_date", "DATE", "Không rỗng ở tầng ứng dụng", "Ngày bắt đầu"),
            ("end_date", "DATE", "end_date ≥ start_date", "Ngày kết thúc"),
            ("status", "VARCHAR", "PENDING / APPROVED / REJECTED", "Trạng thái xử lý"),
            ("admin_note", "TEXT", "Có thể rỗng", "Ghi chú duyệt/từ chối"),
            ("created_at", "DATETIME", "Gán khi tạo", "Thời điểm tạo đơn"),
            ("user_id", "BIGINT", "FK → users.id", "Chủ sở hữu đơn"),
        ),
        (1400, 1800, 2900, 3260),
        body_size=8.7,
        center_columns={0, 1},
    )
    add_caption(doc, "Bảng 4.2. Cấu trúc bảng leave_requests.")
    add_placeholder_figure(doc, placeholder_leaves, "Vị trí chèn ảnh cấu trúc hoặc dữ liệu bảng leave_requests", "Hình 4.3. Cấu trúc hoặc dữ liệu mẫu của bảng leave_requests.")

    add_heading(doc, "4.4. Quy tắc toàn vẹn và trạng thái", 2)
    for item in (
        "Mỗi đơn gắn với một user thông qua user_id.",
        "Đơn mới luôn có trạng thái PENDING.",
        "Chỉ đơn PENDING mới được duyệt, từ chối hoặc hủy.",
        "Nhân viên chỉ được hủy và xem đơn thuộc tài khoản của mình.",
        "Ngày kết thúc không được trước ngày bắt đầu.",
        "Khi từ chối, hệ thống lưu admin_note; khi ghi chú trống, dùng thông báo mặc định.",
    ):
        add_bullet(doc, item, bullet_num_id)
    # Chapter 5.
    add_heading(doc, "CHƯƠNG 5. CÀI ĐẶT CÁC CHỨC NĂNG", 1)
    add_heading(doc, "5.1. Đăng ký và đăng nhập", 2)
    add_body(doc, "AuthController cung cấp các endpoint /register, /login và /logout. Khi đăng ký, controller chuẩn hóa username và họ tên, kiểm tra độ dài, kiểm tra trùng username rồi gọi UserService.save. Service gán vai trò mặc định EMPLOYEE và băm mật khẩu trước khi lưu.")
    add_body(doc, "Khi đăng nhập thành công, đối tượng người dùng được lưu vào session dưới khóa currentUser. Quản trị viên được chuyển đến /admin/dashboard, còn nhân viên được chuyển đến /user/home. Khi đăng xuất, toàn bộ session bị hủy.")
    add_heading(doc, "5.2. Chức năng dành cho nhân viên", 2)
    for item in (
        "Tạo đơn nghỉ phép với lý do và khoảng ngày hợp lệ.",
        "Xem danh sách đơn của bản thân.",
        "Xem chi tiết trạng thái và ghi chú quản trị viên.",
        "Hủy đơn của chính mình khi đơn còn PENDING.",
        "Sử dụng chatbot để tra cứu và nhận hướng dẫn.",
    ):
        add_bullet(doc, item, bullet_num_id)
    add_heading(doc, "5.3. Chức năng dành cho quản trị viên", 2)
    for item in (
        "Dashboard hiển thị tổng số đơn và số lượng theo từng trạng thái.",
        "Danh sách đơn có bộ lọc ALL, PENDING, APPROVED và REJECTED.",
        "Duyệt đơn bằng cách chuyển sang APPROVED; từ chối bằng cách lưu REJECTED và admin_note.",
        "Tìm kiếm nhân viên theo họ tên, username hoặc email.",
        "Cập nhật vai trò nhưng ngăn quản trị viên tự thay đổi quyền của mình.",
    ):
        add_bullet(doc, item, bullet_num_id)
    add_heading(doc, "5.4. Chatbot hỗ trợ tra cứu", 2)
    add_body(doc, "Chatbot hiện được cài đặt theo hướng rule-based và context-aware. ChatbotService chuẩn hóa câu hỏi bằng cách loại dấu tiếng Việt, chuyển về chữ thường và loại ký tự đặc biệt. Sau đó service kiểm tra các nhóm từ khóa, trích xuất mã đơn bằng biểu thức chính quy và gọi LeaveRequestService để lấy dữ liệu.")
    add_body(doc, "Trợ lý hỗ trợ các ý định chính: chào hỏi; xác định vai trò; hướng dẫn tạo hoặc hủy đơn; giải thích trạng thái; tra cứu theo mã; lấy đơn gần nhất; thống kê số lượng; tính tổng ngày nghỉ đã duyệt; và gợi ý thao tác tiếp theo. Nếu câu hỏi vượt quá 500 ký tự, ChatbotController trả về lỗi yêu cầu rút gọn.")
    add_callout(doc, "Bảo vệ dữ liệu", "Khi nhân viên hỏi một mã đơn không thuộc tài khoản, chatbot từ chối hiển thị nội dung. Quản trị viên được phép tra cứu đơn bất kỳ và thấy tên nhân viên tương ứng.")
    add_heading(doc, "5.5. Giao diện chatbot", 2)
    add_body(doc, "Widget chatbot được viết bằng JavaScript và nhúng vào các trang qua component dùng chung. Hội thoại được lưu trong sessionStorage theo userId, giới hạn 30 tin nhắn gần nhất. Giao diện có nút mở/đóng, xóa hội thoại, thao tác nhanh, trạng thái đang trả lời và xử lý lỗi từ endpoint /chatbot/message.")
    add_heading(doc, "5.6. Một số câu hỏi hỗ trợ", 2)
    for item in (
        "Đơn gần nhất của tôi thế nào?",
        "Kiểm tra đơn số 12.",
        "Tôi có bao nhiêu đơn đang chờ duyệt?",
        "Tôi đã nghỉ bao nhiêu ngày trong năm nay?",
        "PENDING nghĩa là gì?",
        "Tôi nên làm gì tiếp theo?",
    ):
        add_bullet(doc, item, bullet_num_id)
    add_heading(doc, "5.7. Luồng xử lý một câu hỏi chatbot", 2)
    chatbot_num_id = add_numbering_definition(doc, bullet=False)
    for item in (
        "Giao diện gửi nội dung câu hỏi đến endpoint /chatbot/message.",
        "Controller kiểm tra session, nội dung rỗng và giới hạn tối đa 500 ký tự.",
        "ChatbotService chuẩn hóa văn bản, nhận diện ý định và trích xuất mã đơn nếu có.",
        "Service nghiệp vụ truy vấn dữ liệu trong đúng phạm vi userId và role.",
        "Câu trả lời được trả về trình duyệt và hội thoại được lưu trong sessionStorage.",
    ):
        add_numbered(doc, item, chatbot_num_id)
    add_heading(doc, "5.8. Nguyên tắc thiết kế trợ lý", 2)
    for item in (
        "Ưu tiên câu trả lời xác định được từ dữ liệu và quy tắc nghiệp vụ hiện có.",
        "Không để mô hình giao diện tự quyết định quyền truy cập hoặc phạm vi dữ liệu.",
        "Khi không nhận diện được ý định, trả hướng dẫn an toàn thay vì suy đoán thông tin.",
    ):
        add_bullet(doc, item, bullet_num_id)

    # Chapter 6.
    add_heading(doc, "CHƯƠNG 6. KIỂM THỬ VÀ ĐÁNH GIÁ", 1)
    add_heading(doc, "6.1. Chiến lược kiểm thử", 2)
    add_body(doc, "Kiểm thử tập trung vào các nghiệp vụ có rủi ro cao: bảo vệ mật khẩu, giới hạn phạm vi dữ liệu chatbot, nhận diện câu hỏi tiếng Việt có hoặc không dấu, tra cứu đơn theo mã, tính số ngày nghỉ và phân biệt các trạng thái dễ nhầm lẫn.")
    add_body(doc, "Bộ kiểm thử JUnit gồm 12 trường hợp: 2 trường hợp cho PasswordService và 10 trường hợp cho ChatbotService. Báo cáo Surefire hiện tại ghi nhận 12/12 kiểm thử đạt, không có lỗi hoặc thất bại.")
    add_heading(doc, "6.2. Các trường hợp kiểm thử tiêu biểu", 2)
    add_styled_table(
        doc,
        ("Mã", "Nội dung", "Kết quả mong đợi", "Kết quả"),
        (
            ("TC01", "Băm và xác thực mật khẩu", "Mật khẩu không lưu rõ; đúng thì khớp, sai thì từ chối", "Đạt"),
            ("TC02", "Tài khoản cũ lưu plain text", "Chỉ chấp nhận khi chuỗi nhập khớp", "Đạt"),
            ("TC03", "Nhân viên hỏi thống kê", "Chỉ trả số liệu cá nhân", "Đạt"),
            ("TC04", "Admin hỏi thống kê", "Trả số liệu toàn hệ thống", "Đạt"),
            ("TC05", "Câu hỏi PENDING không dấu", "Nhận diện đúng trạng thái", "Đạt"),
            ("TC06", "Nhân viên tra cứu đơn người khác", "Từ chối hiển thị dữ liệu", "Đạt"),
            ("TC07", "Tính ngày nghỉ đã duyệt", "Tính bao gồm cả ngày bắt đầu và kết thúc", "Đạt"),
            ("TC08", "Lấy đơn gần nhất", "Chọn bản ghi có createdAt mới nhất", "Đạt"),
        ),
        (900, 2550, 4210, 1700),
        body_size=8.3,
        center_columns={0, 3},
    )
    add_caption(doc, "Bảng 6.1. Các trường hợp kiểm thử tiêu biểu.")
    add_heading(doc, "6.3. Kết quả đạt được", 2)
    for item in (
        "Hoàn thành luồng đăng ký, đăng nhập và phân quyền cơ bản.",
        "Hoàn thành nghiệp vụ tạo, xem, hủy, duyệt và từ chối đơn.",
        "Hoàn thành dashboard, bộ lọc trạng thái và quản lý nhân viên.",
        "Hoàn thành chatbot tra cứu theo vai trò, có lịch sử hội thoại phía trình duyệt.",
        "Áp dụng PBKDF2 cho mật khẩu và có kiểm thử tự động cho các dịch vụ quan trọng.",
    ):
        add_bullet(doc, item, bullet_num_id)
    add_heading(doc, "6.4. Hạn chế", 2)
    for item in (
        "Chưa lưu hạn mức phép năm, ngày lễ hoặc chính sách nghỉ theo phòng ban.",
        "Chatbot sử dụng luật và từ khóa nên chưa hiểu tốt các câu hỏi quá tự do hoặc nhiều ý định.",
        "Chưa tích hợp Spring Security, CSRF token, phân quyền ở tầng cấu hình và cơ chế quên mật khẩu.",
        "Cấu hình database hiện phù hợp phát triển cục bộ, chưa sẵn sàng cho production.",
        "Chưa có thông báo email, nhật ký kiểm toán và quy trình phê duyệt nhiều cấp.",
    ):
        add_bullet(doc, item, bullet_num_id)
    add_heading(doc, "6.5. Đánh giá tổng hợp", 2)
    add_body(doc, "Kết quả kiểm thử cho thấy hệ thống đáp ứng tốt phạm vi học tập và minh họa đúng quy trình nghiệp vụ cốt lõi. Mức độ sẵn sàng được đánh giá theo các tiêu chí sau:")
    add_styled_table(
        doc,
        ("Tiêu chí", "Đánh giá"),
        (
            ("Chức năng", "Hoàn thành các luồng chính cho nhân viên, quản trị viên và chatbot."),
            ("Bảo mật", "Có PBKDF2 và giới hạn dữ liệu theo vai trò; cần bổ sung Spring Security và CSRF."),
            ("Khả dụng", "Giao diện theo vai trò, trạng thái rõ ràng và trợ lý xuất hiện xuyên suốt."),
            ("Bảo trì", "Kiến trúc phân lớp giúp tách controller, service, repository, entity và view."),
            ("Triển khai", "Phù hợp môi trường học tập/cục bộ; chưa nên dùng trực tiếp cho production."),
        ),
        (1800, 7560),
        body_size=9.2,
    )
    add_caption(doc, "Bảng 6.2. Đánh giá tổng hợp mức độ hoàn thiện của hệ thống.")

    # Chapter 7.
    add_heading(doc, "CHƯƠNG 7. HƯỚNG PHÁT TRIỂN", 1)
    add_heading(doc, "7.1. Hoàn thiện nghiệp vụ", 2)
    for item in (
        "Bổ sung hạn mức phép năm, số dư phép và quy tắc trừ phép.",
        "Hỗ trợ nhiều loại nghỉ, ngày lễ, nghỉ nửa ngày và phê duyệt nhiều cấp.",
        "Bổ sung thông báo email hoặc thông báo trong hệ thống khi trạng thái thay đổi.",
        "Tạo lịch nghỉ chung và báo cáo theo phòng ban, tháng hoặc năm.",
    ):
        add_bullet(doc, item, bullet_num_id)
    add_heading(doc, "7.2. Nâng cấp kỹ thuật", 2)
    for item in (
        "Áp dụng Spring Security, CSRF protection, session timeout và audit log.",
        "Chuyển bí mật kết nối database sang biến môi trường và sử dụng connection pool.",
        "Bổ sung REST API, kiểm thử tích hợp và pipeline CI/CD.",
        "Triển khai HTTPS, giám sát ứng dụng và cơ sở dữ liệu được quản lý.",
    ):
        add_bullet(doc, item, bullet_num_id)
    add_heading(doc, "7.3. Phát triển trợ lý AI", 2)
    add_body(doc, "Phiên bản tiếp theo có thể kết hợp bộ luật hiện tại với một mô hình ngôn ngữ lớn để hiểu câu hỏi tự nhiên tốt hơn. Tuy nhiên, mô hình chỉ nên được cấp dữ liệu đã qua lớp kiểm tra quyền; mọi truy vấn đơn vẫn phải được xác thực theo userId và role ở phía server.")
    for item in (
        "Dùng kiến trúc hybrid: intent an toàn và truy vấn dữ liệu do service nội bộ xử lý; LLM chỉ hỗ trợ diễn giải.",
        "Xây dựng knowledge base về quy định nghỉ phép và câu hỏi thường gặp.",
        "Ẩn thông tin nhạy cảm trước khi gửi ngữ cảnh đến dịch vụ AI bên ngoài.",
        "Ghi log, đánh giá câu trả lời và thiết lập giới hạn tốc độ.",
    ):
        add_bullet(doc, item, bullet_num_id)

    # Conclusion.
    add_heading(doc, "KẾT LUẬN", 1)
    add_body(doc, "Đề tài đã xây dựng được hệ thống LeaveFlow với đầy đủ luồng chính cho nhân viên và quản trị viên. Hệ thống tập trung hóa dữ liệu nghỉ phép, làm rõ trạng thái xử lý và giảm các thao tác trao đổi thủ công. Kiến trúc Spring MVC kết hợp service, repository, JPA/Hibernate và JSP giúp mã nguồn có cấu trúc rõ ràng, thuận tiện cho việc mở rộng.")
    add_body(doc, "Chatbot là điểm nhấn của sản phẩm. Dù phiên bản hiện tại hoạt động theo luật, trợ lý đã sử dụng ngữ cảnh vai trò, hỗ trợ tra cứu dữ liệu thực và bảo vệ phạm vi thông tin. Đây là nền tảng phù hợp để phát triển thành trợ lý AI nâng cao trong tương lai mà vẫn duy trì kiểm soát dữ liệu ở phía server.")
    add_body(doc, "Thông qua quá trình thực hiện, sinh viên đã củng cố kiến thức về Java Web, Spring MVC, cơ sở dữ liệu quan hệ, bảo mật mật khẩu, thiết kế giao diện và kiểm thử dịch vụ. Các mục tiêu chính của đề cương đã được hoàn thành; những hạn chế hiện tại đã được xác định rõ để tiếp tục cải tiến.")

    # References.
    add_heading(doc, "TÀI LIỆU THAM KHẢO", 1)
    refs = (
        "Spring Team, Spring Framework Reference Documentation 6.1, https://docs.spring.io/spring-framework/reference/.",
        "Spring Team, Spring Data JPA Reference Documentation, https://docs.spring.io/spring-data/jpa/reference/.",
        "Hibernate Team, Hibernate ORM 6.5 User Guide, https://docs.jboss.org/hibernate/orm/6.5/userguide/html_single/Hibernate_User_Guide.html.",
        "Oracle, Java Platform Standard Edition 17 Documentation, https://docs.oracle.com/en/java/javase/17/.",
        "Oracle, MySQL 8.4 Reference Manual, https://dev.mysql.com/doc/refman/8.4/en/.",
        "Mã nguồn nội bộ dự án LeaveFlow: controller, service, repository, entity, JSP, JavaScript và cấu hình Maven.",
        "Đề cương thực tập thực tế của sinh viên Lê Duy Khang, tháng 7/2026.",
    )
    references_num_id = add_numbering_definition(doc, bullet=False)
    for ref in refs:
        add_numbered(doc, ref, references_num_id)

    # Appendix A.
    add_heading(doc, "PHỤ LỤC A. HƯỚNG DẪN CÀI ĐẶT VÀ CHẠY", 1)
    add_heading(doc, "A.1. Yêu cầu môi trường", 2)
    for item in (
        "JDK 17.",
        "Apache Maven tương thích với Java 17.",
        "Apache Tomcat hỗ trợ Jakarta Servlet 6.",
        "MySQL 8.x và database tên leave_management.",
    ):
        add_bullet(doc, item, bullet_num_id)
    add_heading(doc, "A.2. Các bước cài đặt", 2)
    appendix_a_num_id = add_numbering_definition(doc, bullet=False)
    for item in (
        "Tạo database leave_management trong MySQL.",
        "Cập nhật jdbc.url, jdbc.username và jdbc.password trong jdbc.properties.",
        "Chạy lệnh mvn clean package để tạo leave_management.war.",
        "Chép file WAR vào thư mục webapps của Tomcat và khởi động server.",
        "Truy cập ứng dụng theo context path của WAR, đăng ký tài khoản nhân viên và cấu hình một tài khoản ADMIN trong database khi cần kiểm thử quản trị.",
    ):
        add_numbered(doc, item, appendix_a_num_id)
    add_callout(doc, "Khuyến nghị", "Không sử dụng tài khoản MySQL root hoặc mật khẩu trống khi triển khai ngoài máy cá nhân. Nên dùng biến môi trường và tài khoản database có quyền tối thiểu.")

    # Appendix B.
    add_heading(doc, "PHỤ LỤC B. HƯỚNG DẪN THAY ẢNH DATABASE", 1)
    add_body(doc, "Ba ảnh placeholder nằm trong Chương 4 đã được chèn dưới dạng hình ảnh Word với kích thước cố định. Thay ảnh theo các bước sau để giữ nguyên bố cục:")
    appendix_b_num_id = add_numbering_definition(doc, bullet=False)
    for item in (
        "Chuẩn bị ảnh rõ nét, ưu tiên tỷ lệ ngang và độ phân giải từ 1200 px trở lên.",
        "Nhấp chuột phải vào placeholder tương ứng.",
        "Chọn Change Picture → From a File.",
        "Chọn ảnh ERD, bảng users hoặc bảng leave_requests.",
        "Kiểm tra phần Crop nếu ảnh bị cắt; giữ nguyên chú thích Hình 4.1, 4.2 hoặc 4.3.",
        "Mở mục lục và chọn Update Field → Update entire table trước khi nộp.",
    ):
        add_numbered(doc, item, appendix_b_num_id)
    add_callout(doc, "Gợi ý ảnh", "Có thể dùng ảnh ERD từ MySQL Workbench hoặc ảnh chụp phần Structure/Data của phpMyAdmin/MySQL Workbench. Hãy che mật khẩu hoặc dữ liệu nhạy cảm trước khi chèn.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build_document()
